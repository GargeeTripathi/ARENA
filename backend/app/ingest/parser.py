"""pdf / docx / txt / md -> pages of text (tables rendered as markdown, flagged)."""
from dataclasses import dataclass, field
from pathlib import Path


class ParseError(ValueError):
    pass


@dataclass
class Page:
    page: int  # 1-based
    text: str
    tables: list[str] = field(default_factory=list)  # each table as markdown


def _table_md(rows: list[list]) -> str:
    rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in rows if r and any(r)]
    if not rows:
        return ""
    head, *body = rows
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _pdf(path: Path) -> list[Page]:
    import pymupdf
    import pdfplumber

    pages: list[Page] = []
    with pymupdf.open(path) as doc, pdfplumber.open(path) as plumb:
        for i, p in enumerate(doc):
            tables = []
            try:
                tables = [t for t in (_table_md(tb) for tb in plumb.pages[i].extract_tables()) if t]
            except Exception:  # table extraction is best-effort
                pass
            pages.append(Page(i + 1, p.get_text("text"), tables))
    return pages


def _docx(path: Path) -> list[Page]:
    import docx

    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    tables = [_table_md([[c.text for c in r.cells] for r in t.rows]) for t in d.tables]
    return [Page(1, text, [t for t in tables if t])]


def parse(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages = _pdf(path)
    elif ext == ".docx":
        pages = _docx(path)
    elif ext in (".txt", ".md"):
        pages = [Page(1, path.read_text(encoding="utf-8", errors="replace"))]
    else:
        raise ParseError(f"{path.name}: unsupported file type {ext}")
    if not any(p.text.strip() or p.tables for p in pages):
        raise ParseError(f"{path.name}: no extractable text (empty or scanned document; OCR not supported)")
    return pages
